"""
General Agent 工具 — 知识库检索 / 通用问答

用于处理非检测类的一般性问题，如概念解释、使用帮助等。
search_knowledge 接入了真正的 RAG 向量检索（knowledge_service + embedding_service），
检索范围按调用方 user_id 做了权限边界过滤（自己的文档 + 系统内置文档）。
"""

import io
import json

from langchain_core.tools import tool

from app.core.logger import get_logger

logger = get_logger(__name__)


def _make_general_tools(user_id: int):
    """创建绑定到当前用户的 General Agent 工具"""

    @tool
    def search_knowledge(query: str, top_k: int = 5) -> str:
        """
        搜索知识库。用于回答用户关于缺陷类型、检测原理、模型参数等一般性问题。
        只会检索当前用户可见的知识文档（自己上传的 + 系统内置的资料）。

        Args:
            query: 搜索关键词或问题
            top_k: 返回最相关的分块数量，默认 5

        Returns:
            JSON 字符串，包含检索到的分块列表（chunk 原文 + 所属文档标题）；
            知识库暂无匹配内容或未配置 embedding 服务时，返回 found=false。
        """
        from app.database.session import SessionLocal
        from app.services.embedding_service import embedding_service
        from app.services.knowledge_service import knowledge_service

        db = SessionLocal()
        try:
            query_embedding = embedding_service.embed_text(query)
            chunks = knowledge_service.search_similar_chunks(
                db,
                query_embedding=query_embedding,
                user_id=user_id,
                is_superuser=False,
                top_k=top_k,
            )
            if not chunks:
                return json.dumps({"found": False, "message": "知识库中暂无相关内容"}, ensure_ascii=False)

            results = [
                {
                    "document_id": c.document.id if c.document else None,
                    "chunk_id": c.id,
                    "document_title": c.document.title if c.document else "",
                    "content": c.content,
                }
                for c in chunks
            ]
            return json.dumps({"found": True, "results": results}, ensure_ascii=False)
        except Exception as exc:
            logger.warning("search_knowledge 工具异常: %s", exc)
            return json.dumps(
                {"found": False, "message": "知识库检索暂时不可用"}, ensure_ascii=False
            )
        finally:
            db.close()

    @tool
    def read_attachment(path: str) -> str:
        """Read common user-provided text, office and PDF attachments safely."""
        from pathlib import Path
        from app.api.agent import _validate_uploaded_image_path
        from app.services.document_parser import extract_text

        safe_path = _validate_uploaded_image_path(path, user_id=user_id)
        filename = Path(safe_path).name
        suffix = Path(filename).suffix.lower()
        raw = Path(safe_path).read_bytes()
        if suffix in {".txt", ".md", ".csv", ".json", ".yaml", ".yml", ".log", ".xml"}:
            content = raw.decode("utf-8", errors="replace")
        elif suffix in {".pdf", ".docx"}:
            content = extract_text(filename, raw)
        elif suffix == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            rows = []
            for sheet in workbook.worksheets:
                rows.append(f"[{sheet.title}]")
                rows.extend("\t".join("" if value is None else str(value) for value in row) for row in sheet.iter_rows(values_only=True))
            content = "\n".join(rows)
        elif suffix == ".pptx":
            from pptx import Presentation
            presentation = Presentation(io.BytesIO(raw))
            content = "\n".join(
                shape.text for slide in presentation.slides for shape in slide.shapes
                if hasattr(shape, "text") and shape.text
            )
        else:
            return json.dumps({"name": filename, "readable": False, "message": "Attachment was uploaded but has no text extractor; use a compatible tool or convert it first."}, ensure_ascii=False)
        return json.dumps({"name": filename, "content": content[:20000]}, ensure_ascii=False)

    return [search_knowledge, read_attachment]
